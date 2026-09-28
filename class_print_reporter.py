# -*- coding: utf-8 -*-
"""Compact A4 class report cards for coach-side batch printing.

The public entry point :func:`create_class_print_report` accepts the same
structured individual-summary payload used by the coach dashboard.  Each
student is rendered as one indivisible card; Word page breaks are inserted
after exactly two or three cards so teachers can print, cut, and distribute
the reports without manually reformatting them.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_X_CM = 0.72
MARGIN_Y_CM = 0.62
CONTENT_WIDTH_CM = PAGE_WIDTH_CM - MARGIN_X_CM * 2

NAVY = "10233F"
TEAL = "0F766E"
TEAL_LIGHT = "DFF7F3"
BLUE_LIGHT = "EAF2FF"
AMBER_LIGHT = "FFF5D6"
SLATE_50 = "F8FAFC"
SLATE_200 = "CBD5E1"
SLATE_500 = "64748B"
SLATE_700 = "334155"
WHITE = "FFFFFF"

RADAR_KEYS = (
    ("approach_rhythm", "助跑节奏"),
    ("support_stability", "支撑稳定"),
    ("backswing_folding", "后摆折叠"),
    ("ankle_rigidity", "脚踝刚性"),
    ("whipping_velocity", "鞭打速度"),
)


def _safe_text(value: Any, fallback: str = "暂无数据") -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text or fallback


def _compact(value: Any, limit: int) -> str:
    text = _safe_text(value)
    if len(text) <= limit:
        return text
    return text[: max(1, limit - 1)].rstrip("，。；;,. ") + "…"


def _safe_float(value: Any) -> Optional[float]:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number or number in (float("inf"), float("-inf")):
        return None
    return number


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, *, color: str = SLATE_200, size: int = 8, style: str = "single") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), style)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_table_width(table, width_cm: float) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tbl_w.set(qn("w:type"), "dxa")


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_run_font(run, size: float, *, bold: bool = False, color: str = SLATE_700) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def _paragraph(cell, text: str = "", *, size: float = 8.0, bold: bool = False,
               color: str = SLATE_700, align=WD_ALIGN_PARAGRAPH.LEFT,
               before: float = 0, after: float = 0, line: float = 1.0):
    paragraph = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    _clear_paragraph(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = line
    run = paragraph.add_run(text)
    _set_run_font(run, size, bold=bold, color=color)
    return paragraph


def _add_label_value(cell, label: str, value: Any, *, size: float, limit: int) -> None:
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    label_run = paragraph.add_run(f"{label}｜")
    _set_run_font(label_run, size, bold=True, color=TEAL)
    value_run = paragraph.add_run(_compact(value, limit))
    _set_run_font(value_run, size, color=SLATE_700)


def _font(size: int, *, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size=size)
            except OSError:
                continue
    return ImageFont.load_default()


def _radar_image(scores: Mapping[str, Any], size: tuple[int, int] = (720, 430)) -> io.BytesIO:
    import math

    width, height = size
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    cx, cy = width // 2, height // 2 + 4
    radius = min(width * 0.27, height * 0.37)
    label_font = _font(max(18, int(height * 0.045)), bold=True)
    value_font = _font(max(17, int(height * 0.042)))
    angles = [-math.pi / 2 + 2 * math.pi * i / 5 for i in range(5)]

    for level in (0.25, 0.5, 0.75, 1.0):
        points = [(cx + radius * level * math.cos(a), cy + radius * level * math.sin(a)) for a in angles]
        draw.polygon(points, outline="#CBD5E1", width=2)
    for angle in angles:
        draw.line((cx, cy, cx + radius * math.cos(angle), cy + radius * math.sin(angle)), fill="#D9E2EC", width=2)

    values = []
    for key, _ in RADAR_KEYS:
        raw = _safe_float(scores.get(key))
        values.append(max(0.0, min(100.0, raw if raw is not None else 0.0)))
    data_points = [
        (cx + radius * value / 100 * math.cos(angle), cy + radius * value / 100 * math.sin(angle))
        for value, angle in zip(values, angles)
    ]
    draw.polygon(data_points, fill="#8ADDD2", outline="#0F766E", width=5)
    for x, y in data_points:
        draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill="#0F766E")

    for (_, label), value, angle in zip(RADAR_KEYS, values, angles):
        x = cx + (radius + 53) * math.cos(angle)
        y = cy + (radius + 38) * math.sin(angle)
        bbox = draw.textbbox((0, 0), label, font=label_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, y - 14), label, font=label_font, fill="#334155")
        value_text = f"{value:.0f}"
        vb = draw.textbbox((0, 0), value_text, font=value_font)
        draw.text((x - (vb[2] - vb[0]) / 2, y + 12), value_text, font=value_font, fill="#0F766E")

    stream = io.BytesIO()
    image.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    return stream


def _visual_image(raw: Any, size: tuple[int, int] = (720, 430)) -> io.BytesIO:
    width, height = size
    try:
        if isinstance(raw, (bytes, bytearray)):
            source = Image.open(io.BytesIO(raw)).convert("RGB")
        elif isinstance(raw, (str, Path)) and Path(raw).is_file():
            source = Image.open(str(raw)).convert("RGB")
        else:
            raise ValueError("missing")
        fitted = ImageOps.contain(source, (width - 20, height - 20))
        canvas = Image.new("RGB", size, "#081525")
        canvas.paste(fitted, ((width - fitted.width) // 2, (height - fitted.height) // 2))
    except Exception:
        canvas = Image.new("RGB", size, "#EEF2F6")
        draw = ImageDraw.Draw(canvas)
        draw.rounded_rectangle((14, 14, width - 14, height - 14), radius=24, outline="#94A3B8", width=4)
        cx, cy = width // 2, height // 2 - 30
        draw.ellipse((cx - 28, cy - 86, cx + 28, cy - 30), outline="#64748B", width=7)
        draw.line((cx, cy - 28, cx, cy + 66), fill="#64748B", width=8)
        draw.line((cx, cy + 2, cx - 72, cy + 28), fill="#64748B", width=8)
        draw.line((cx, cy + 2, cx + 75, cy - 5), fill="#64748B", width=8)
        draw.line((cx, cy + 65, cx - 50, cy + 135), fill="#64748B", width=8)
        draw.line((cx, cy + 65, cx + 78, cy + 108), fill="#64748B", width=8)
        label = "暂无可用击球定格图"
        font = _font(26, bold=True)
        bbox = draw.textbbox((0, 0), label, font=font)
        draw.text(((width - bbox[2] + bbox[0]) / 2, height - 55), label, font=font, fill="#64748B")
    stream = io.BytesIO()
    canvas.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    return stream


def _format_period(report: Mapping[str, Any]) -> str:
    requested = report.get("requestedPeriod") if isinstance(report.get("requestedPeriod"), Mapping) else {}
    actual = report.get("period") if isinstance(report.get("period"), Mapping) else {}
    start = requested.get("start") or actual.get("start") or "最早记录"
    end = requested.get("end") or actual.get("end") or "最新记录"
    return f"{start} 至 {end}"


def _score_line(report: Mapping[str, Any]) -> str:
    summary = report.get("scoreSummary") if isinstance(report.get("scoreSummary"), Mapping) else {}
    mean = _safe_float(summary.get("mean"))
    latest = _safe_float(summary.get("latest"))
    change = _safe_float(summary.get("change"))
    parts = [
        f"均分 {mean:.1f}" if mean is not None else "均分 --",
        f"最近 {latest:.1f}" if latest is not None else "最近 --",
    ]
    if change is not None:
        parts.append(f"变化 {change:+.1f}")
    parts.append(f"A级有效 {int(report.get('formalAttemptCount') or 0)} 次")
    return "  ·  ".join(parts)


def _add_visual(cell, stream: io.BytesIO, caption: str, *, density: int) -> None:
    paragraph = cell.paragraphs[0]
    _clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run()
    run.add_picture(stream, width=Cm(8.7), height=Cm(3.28 if density == 2 else 2.55))
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(0)
    cap_run = cap.add_run(caption)
    _set_run_font(cap_run, 7 if density == 2 else 6.4, bold=True, color=SLATE_500)


def _add_student_card(document: Document, report: Mapping[str, Any], *, density: int,
                      school: str, class_group: str) -> None:
    compact_mode = density == 3
    container = document.add_table(rows=1, cols=1)
    container.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_width(container, CONTENT_WIDTH_CM)
    _set_table_borders(container, color="91A4B8", size=8, style="dashed")
    _prevent_row_split(container.rows[0])
    container_cell = container.cell(0, 0)
    _set_cell_margins(container_cell, top=0, start=0, bottom=0, end=0)
    empty = container_cell.paragraphs[0]
    empty.paragraph_format.space_before = Pt(0)
    empty.paragraph_format.space_after = Pt(0)
    empty.paragraph_format.line_spacing = 0.1
    empty_run = empty.add_run()
    empty_run.font.size = Pt(1)

    card = container_cell.add_table(rows=5, cols=1)
    card.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_width(card, CONTENT_WIDTH_CM)
    _set_table_borders(card, color=SLATE_200, size=5, style="single")
    for row in card.rows:
        _prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=45 if compact_mode else 65, bottom=45 if compact_mode else 65)

    student_id = _safe_text(report.get("studentId"), "未填写编号")
    header = card.cell(0, 0)
    _set_cell_shading(header, NAVY)
    header_p = header.paragraphs[0]
    _clear_paragraph(header_p)
    header_p.paragraph_format.space_before = Pt(0)
    header_p.paragraph_format.space_after = Pt(0)
    left = header_p.add_run(f"个人总体分析  |  {student_id}")
    _set_run_font(left, 10 if compact_mode else 11.5, bold=True, color=WHITE)
    right = header_p.add_run(f"    {school} · {class_group}")
    _set_run_font(right, 7 if compact_mode else 8, color="BFD6E8")

    meta = card.cell(1, 0)
    _set_cell_shading(meta, TEAL_LIGHT)
    meta_p = meta.paragraphs[0]
    _clear_paragraph(meta_p)
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(0)
    meta_run = meta_p.add_run(f"分析时段：{_format_period(report)}    {_score_line(report)}")
    _set_run_font(meta_run, 7.2 if compact_mode else 8.2, bold=True, color=TEAL)

    visuals_cell = card.cell(2, 0)
    visuals = visuals_cell.add_table(rows=1, cols=2)
    visuals.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(visuals, color=SLATE_200, size=5)
    widths = [Cm(CONTENT_WIDTH_CM / 2 - 0.2)] * 2
    for idx, visual_cell in enumerate(visuals.rows[0].cells):
        visual_cell.width = widths[idx]
        _set_cell_margins(visual_cell, top=35, start=35, bottom=25, end=35)
    _add_visual(
        visuals.cell(0, 0),
        _visual_image(report.get("impactFrameBytes")),
        "击球瞬间骨骼定格",
        density=density,
    )
    errors = report.get("topErrors") if isinstance(report.get("topErrors"), list) else []
    error_text = "、".join(str(item.get("label") or "") for item in errors[:2] if isinstance(item, Mapping))
    radar_caption = "五维能力整体画像" + (f" · 高频：{_compact(error_text, 18)}" if error_text else "")
    _add_visual(
        visuals.cell(0, 1),
        _radar_image(report.get("fiveDimensionScores") or {}),
        radar_caption,
        density=density,
    )

    analysis_cell = card.cell(3, 0)
    analysis_table = analysis_cell.add_table(rows=1, cols=2)
    analysis_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(analysis_table, color=SLATE_200, size=5)
    left_cell, right_cell = analysis_table.rows[0].cells
    for item in (left_cell, right_cell):
        _set_cell_margins(item, top=35, start=70, bottom=35, end=70)
        item.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    font_size = 6.8 if compact_mode else 7.7
    small_limit = 52 if compact_mode else 92
    _paragraph(left_cell, "AIGC 纵向结论", size=7.4 if compact_mode else 8.5, bold=True, color=TEAL)
    _add_label_value(left_cell, "总体", report.get("overallAssessment"), size=font_size, limit=small_limit)
    _add_label_value(left_cell, "进步", report.get("progressAnalysis"), size=font_size, limit=small_limit)
    _add_label_value(left_cell, "优势", report.get("strengths"), size=font_size, limit=small_limit)
    _paragraph(right_cell, "AIGC 训练处方", size=7.4 if compact_mode else 8.5, bold=True, color=TEAL)
    _add_label_value(right_cell, "盲区", report.get("weaknesses"), size=font_size, limit=small_limit)
    _add_label_value(right_cell, "练法", report.get("prescription"), size=font_size, limit=small_limit)
    _add_label_value(right_cell, "剂量", report.get("dosage"), size=font_size, limit=small_limit)

    task_cell = card.cell(4, 0)
    _set_cell_shading(task_cell, AMBER_LIGHT)
    task = report.get("selfCheckTask") if isinstance(report.get("selfCheckTask"), Mapping) else {}
    title = _safe_text(task.get("title"), "专属图式自查任务")
    cue = task.get("cue") or task.get("instruction") or task.get("description") or "按处方完成动作并对照关键姿势自查。"
    dosage = task.get("dosage") or report.get("dosage") or "3组×5次"
    task_p = task_cell.paragraphs[0]
    _clear_paragraph(task_p)
    task_p.paragraph_format.space_before = Pt(0)
    task_p.paragraph_format.space_after = Pt(0)
    task_run = task_p.add_run(
        f"图式自查任务｜{_compact(title, 24)}  ·  {_compact(cue, 52 if compact_mode else 80)}  ·  {dosage}    "
        "□ 第1次  □ 第2次  □ 第3次    教练签字：________"
    )
    _set_run_font(task_run, 6.7 if compact_mode else 7.7, bold=True, color="8A4B08")

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1 if compact_mode else 3)
    spacer.paragraph_format.line_spacing = 0.45


def _set_document_defaults(document: Document, *, school: str, class_group: str,
                           period_text: str, generated_at: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(MARGIN_Y_CM)
    section.bottom_margin = Cm(MARGIN_Y_CM)
    section.left_margin = Cm(MARGIN_X_CM)
    section.right_margin = Cm(MARGIN_X_CM)
    section.header_distance = Cm(0.22)
    section.footer_distance = Cm(0.22)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(1)
    hrun = hp.add_run(f"足球AI可视化反馈系统  ·  {school}  ·  {class_group}  ·  {period_text}")
    _set_run_font(hrun, 7.2, bold=True, color=TEAL)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    frun = fp.add_run(f"沿卡片虚线裁剪后发放  ·  生成时间 {generated_at}  ·  学生个人反馈资料")
    _set_run_font(frun, 6.6, color=SLATE_500)


def create_class_print_report(
    reports: Iterable[Mapping[str, Any]],
    *,
    school: str,
    class_group: str,
    period_text: str,
    generated_at: str,
    students_per_page: int,
    output_path: str | Path,
) -> Path:
    """Create an A4 portrait Word report with exactly 2 or 3 student cards/page."""

    density = int(students_per_page)
    if density not in (2, 3):
        raise ValueError("students_per_page must be 2 or 3")
    report_list = [item for item in reports if isinstance(item, Mapping)]
    if not report_list:
        raise ValueError("at least one student report is required")

    document = Document()
    _set_document_defaults(
        document,
        school=_safe_text(school, "未设置学校"),
        class_group=_safe_text(class_group, "未设置班级"),
        period_text=_safe_text(period_text, "全部时段"),
        generated_at=_safe_text(generated_at),
    )

    for index, report in enumerate(report_list):
        if index and index % density == 0:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        _add_student_card(
            document,
            report,
            density=density,
            school=_safe_text(school, "未设置学校"),
            class_group=_safe_text(class_group, "未设置班级"),
        )

    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return target


__all__ = ["create_class_print_report"]
